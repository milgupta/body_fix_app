#!/usr/bin/env python3
import argparse
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


API_BASE = "https://api.cloudflare.com/client/v4"


def load_dotenv(path=Path(".env")):
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def parse_args():
    load_dotenv()
    parser = argparse.ArgumentParser(
        description="Upload GetBodyFix stretch videos to Cloudflare Stream and create stretch_videos.json."
    )
    parser.add_argument(
        "--audit",
        default="/Users/milgupta/Downloads/GetBodyFix_Raw/GetBodyFix_Video_Audit.xlsx",
        help="Path to GetBodyFix_Video_Audit.xlsx.",
    )
    parser.add_argument(
        "--source-dir",
        default="/Users/milgupta/Downloads/GetBodyFix_Raw",
        help="Directory containing the audited video files.",
    )
    parser.add_argument(
        "--output",
        default="BodyFix/Data/stretch_videos.json",
        help="Output JSON path for app video mapping.",
    )
    parser.add_argument(
        "--progress",
        default="upload_progress_cloudflare_stream.json",
        help="Progress JSON path used to resume uploads.",
    )
    parser.add_argument(
        "--account-id",
        default=os.environ.get("CF_ACCOUNT_ID"),
        help="Cloudflare account ID. Defaults to CF_ACCOUNT_ID env var.",
    )
    parser.add_argument(
        "--api-token",
        default=os.environ.get("CF_API_TOKEN"),
        help="Cloudflare API token. Defaults to CF_API_TOKEN env var.",
    )
    parser.add_argument(
        "--customer-code",
        default=os.environ.get("CF_CUSTOMER_CODE"),
        help="Optional Cloudflare Stream customer code, the part after customer-. Defaults to CF_CUSTOMER_CODE env var.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Upload only the first N rows. Useful for testing.",
    )
    parser.add_argument(
        "--poll",
        action="store_true",
        help="Poll each uploaded video until Cloudflare reports readyToStream.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Do not upload; print the rows that would be uploaded.",
    )
    parser.add_argument(
        "--import-existing",
        action="store_true",
        help="Do not upload; build the output JSON from videos already in Cloudflare Stream by matching filenames.",
    )
    return parser.parse_args()


def require_openpyxl():
    if load_workbook is None:
        raise SystemExit(
            "Missing openpyxl. Run with the bundled Python runtime or install openpyxl for your python3."
        )


def load_audit_rows(audit_path, source_dir):
    require_openpyxl()
    workbook = load_workbook(audit_path, read_only=True, data_only=True)
    sheet = workbook["Video Audit"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    indexes = {header: index for index, header in enumerate(headers)}
    required = ["Video file", "Matched stretch", "Stretch ID"]
    missing_headers = [header for header in required if header not in indexes]
    if missing_headers:
        raise SystemExit(f"Missing expected workbook columns: {', '.join(missing_headers)}")

    rows = []
    seen_stretch_ids = set()
    duplicate_stretch_ids = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        filename = row[indexes["Video file"]]
        stretch_id = row[indexes["Stretch ID"]]
        name = row[indexes["Matched stretch"]]
        if not filename or not stretch_id:
            continue
        if stretch_id in seen_stretch_ids:
            duplicate_stretch_ids.append((stretch_id, filename))
            continue
        seen_stretch_ids.add(stretch_id)
        video_path = resolve_video_path(source_dir, filename)
        rows.append(
            {
                "stretch_id": stretch_id,
                "name": name,
                "filename": filename,
                "actual_filename": video_path.name,
                "path": video_path,
            }
        )
    if duplicate_stretch_ids:
        print("Skipping duplicate stretch IDs from audit sheet:")
        for stretch_id, filename in duplicate_stretch_ids:
            print(f"  {stretch_id}: {filename}")
    return rows


def resolve_video_path(source_dir, filename):
    video_path = source_dir / filename
    if video_path.exists():
        return video_path

    stem = Path(filename).stem
    suffix = Path(filename).suffix
    cleaned_stem = re.sub(r"\s*\([^)]*\)\s*", "", stem).strip()
    cleaned_path = source_dir / f"{cleaned_stem}{suffix}"
    if cleaned_path.exists():
        return cleaned_path

    return video_path


def load_progress(path):
    if not path.exists():
        return {"uploads": {}}
    with path.open() as file:
        return json.load(file)


def save_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_suffix(path.suffix + ".tmp")
    with temp_path.open("w") as file:
        json.dump(value, file, indent=2, sort_keys=True)
        file.write("\n")
    temp_path.replace(path)


def run_curl_upload(account_id, api_token, video_path):
    url = f"{API_BASE}/accounts/{account_id}/stream"
    command = [
        "curl",
        "--request",
        "POST",
        "--silent",
        "--show-error",
        "--fail-with-body",
        "--header",
        f"Authorization: Bearer {api_token}",
        "--form",
        f"file=@{video_path}",
        url,
    ]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(
            f"Cloudflare upload failed for {video_path.name}:\n{result.stderr}\n{result.stdout}"
        )
    payload = json.loads(result.stdout)
    if not payload.get("success"):
        raise RuntimeError(f"Cloudflare upload failed for {video_path.name}: {payload}")
    return payload["result"]


def fetch_video_details(account_id, api_token, uid):
    url = f"{API_BASE}/accounts/{account_id}/stream/{uid}"
    command = [
        "curl",
        "--silent",
        "--show-error",
        "--fail-with-body",
        "--header",
        f"Authorization: Bearer {api_token}",
        url,
    ]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"Cloudflare details request failed for {uid}:\n{result.stderr}\n{result.stdout}")
    payload = json.loads(result.stdout)
    if not payload.get("success"):
        raise RuntimeError(f"Cloudflare details request failed for {uid}: {payload}")
    return payload["result"]


def list_existing_videos(account_id, api_token):
    videos = []
    page = 1
    while True:
        url = f"{API_BASE}/accounts/{account_id}/stream?per_page=1000&page={page}"
        command = [
            "curl",
            "--silent",
            "--show-error",
            "--fail-with-body",
            "--header",
            f"Authorization: Bearer {api_token}",
            url,
        ]
        result = subprocess.run(command, capture_output=True, text=True)
        if result.returncode != 0:
            raise RuntimeError(f"Cloudflare list request failed:\n{result.stderr}\n{result.stdout}")
        payload = json.loads(result.stdout)
        if not payload.get("success"):
            raise RuntimeError(f"Cloudflare list request failed: {payload}")
        batch = payload.get("result", [])
        videos.extend(batch)
        info = payload.get("result_info") or {}
        total_pages = info.get("total_pages")
        if not batch or total_pages is None or page >= total_pages:
            break
        page += 1
    return videos


def wait_until_ready(account_id, api_token, uid, timeout_seconds=300):
    started = time.time()
    while True:
        details = fetch_video_details(account_id, api_token, uid)
        if details.get("readyToStream"):
            return details
        if time.time() - started > timeout_seconds:
            return details
        time.sleep(5)


def hls_url(customer_code, uid):
    return f"https://customer-{customer_code}.cloudflarestream.com/{uid}/manifest/video.m3u8"


def build_output(progress, rows, customer_code):
    videos = {}
    for row in rows:
        upload = progress["uploads"].get(row["stretch_id"])
        if not upload:
            continue
        uid = upload["uid"]
        video = {
            "uid": uid,
            "filename": row["actual_filename"],
            "name": row["name"],
        }
        hls = upload.get("hlsURL")
        if not hls and customer_code:
            hls = hls_url(customer_code, uid)
        if hls:
            video["hlsURL"] = hls
        videos[row["stretch_id"]] = video
    return {
        "provider": "cloudflare_stream",
        "customerCode": customer_code,
        "videos": videos,
    }


def import_existing_videos(account_id, api_token, progress, rows):
    videos = list_existing_videos(account_id, api_token)
    by_name = {}
    for video in videos:
        name = (video.get("meta", {}).get("name") or video.get("filename") or "").strip()
        if not name:
            continue
        existing = by_name.get(name)
        if existing is None or (video.get("created") or "") > (existing.get("created") or ""):
            by_name[name] = video

    missing = []
    for row in rows:
        video = by_name.get(row["actual_filename"])
        if video is None:
            missing.append((row["stretch_id"], row["actual_filename"]))
            continue
        uid = video["uid"]
        progress["uploads"][row["stretch_id"]] = {
            "uid": uid,
            "filename": row["actual_filename"],
            "name": row["name"],
            "readyToStream": bool(video.get("readyToStream")),
            "hlsURL": video.get("playback", {}).get("hls"),
        }

    return missing


def main():
    args = parse_args()
    audit_path = Path(args.audit)
    source_dir = Path(args.source_dir)
    output_path = Path(args.output)
    progress_path = Path(args.progress)

    rows = load_audit_rows(audit_path, source_dir)
    if args.limit is not None:
        rows = rows[: args.limit]

    missing_files = [str(row["path"]) for row in rows if not row["path"].exists()]
    if missing_files:
        raise SystemExit("Missing video files:\n" + "\n".join(missing_files))

    if args.dry_run:
        print(f"Would upload {len(rows)} videos.")
        for row in rows:
            if row["filename"] == row["actual_filename"]:
                print(f"{row['stretch_id']}: {row['actual_filename']}")
            else:
                print(f"{row['stretch_id']}: {row['actual_filename']} (resolved from {row['filename']})")
        return

    if not args.account_id:
        raise SystemExit("Missing Cloudflare account ID. Set CF_ACCOUNT_ID or pass --account-id.")
    if not args.api_token:
        raise SystemExit("Missing Cloudflare API token. Set CF_API_TOKEN or pass --api-token.")

    progress = load_progress(progress_path)
    progress.setdefault("uploads", {})

    if args.import_existing:
        missing = import_existing_videos(args.account_id, args.api_token, progress, rows)
        save_json(progress_path, progress)
        output = build_output(progress, rows, args.customer_code)
        save_json(output_path, output)
        print(f"Imported {len(output['videos'])} existing video mappings to {output_path}")
        if missing:
            print("Missing existing Cloudflare videos for:")
            for stretch_id, filename in missing:
                print(f"  {stretch_id}: {filename}")
            raise SystemExit(1)
        return

    for index, row in enumerate(rows, start=1):
        stretch_id = row["stretch_id"]
        if stretch_id in progress["uploads"]:
            print(f"[{index}/{len(rows)}] Skipping {stretch_id}; already uploaded.")
            continue

        print(f"[{index}/{len(rows)}] Uploading {stretch_id} from {row['filename']}...")
        result = run_curl_upload(args.account_id, args.api_token, row["path"])
        uid = result["uid"]
        details = wait_until_ready(args.account_id, args.api_token, uid) if args.poll else result
        progress["uploads"][stretch_id] = {
            "uid": uid,
            "filename": row["actual_filename"],
            "name": row["name"],
            "readyToStream": bool(details.get("readyToStream")),
            "hlsURL": details.get("playback", {}).get("hls"),
        }
        save_json(progress_path, progress)

    output = build_output(progress, rows, args.customer_code)
    save_json(output_path, output)
    print(f"Saved {len(output['videos'])} video mappings to {output_path}")


if __name__ == "__main__":
    main()
